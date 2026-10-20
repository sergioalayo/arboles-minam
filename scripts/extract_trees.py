"""Fase 1 — Genera `data/arboles.json` a partir de los diagramas de `raw-docs/`.

Flujo por archivo:
1. Formas con geometría absoluta (`shapes.py`); los Word se leen desde la copia normalizada
   por `normalize_word.py` (anclas a página, .doc convertido a DrawingML).
2. Detección del problema central y separación en semiplanos (efectos arriba, causas abajo).
3. Relaciones padre-hijo, en orden de prioridad:
   a. Códigos jerárquicos explícitos ("CI 2.3" → "CD 2", "1.4" → "1").
   b. Conectores: componentes conexos de líneas que tocan cajas; dentro de cada componente,
      la(s) caja(s) más cercana(s) al problema central es (son) el padre.
   c. Geometría: caja más cercana hacia el centro, con distinto estilo y solape horizontal.
4. Niveles (camino más largo desde el problema central), tipos, IDs semánticos y metadatos.

Uso:  python scripts/extract_trees.py [--report ruta.md]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import docx  # python-docx
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from classify import (ORGANOS, TRANSVERSAL, assign_severity, classify_eje, classify_organo,  # noqa: E402
                      extract_evidencia, make_label)
from shapes import Shape, docx_shapes, pptx_slides  # noqa: E402
from shapes import _paragraph_text as paragraph_text  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / 'raw-docs'
NORMALIZED = Path(__file__).resolve().parent / '.work' / 'normalized'
OUT = ROOT / 'data' / 'arboles.json'
OUT_ORGANOS = ROOT / 'data' / 'organos.json'

CM = 360000  # EMU por centímetro
TOUCH = 0.22 * CM
ROW_TOL = 0.5 * CM

# Metadatos por mesa: tema (eje del problema central) y órgano por defecto.
MESAS: dict[str, dict[str, str]] = {
    '1': {'tema': 'Gobernanza ambiental', 'organo': TRANSVERSAL},
    '2': {'tema': 'Simplificación administrativa', 'organo': 'SG'},
    '3': {'tema': 'Ciencia e información ambiental', 'organo': TRANSVERSAL},
    '4': {'tema': 'Brechas sectoriales', 'organo': 'OGPP'},
    '5': {'tema': 'Programas, proyectos especiales y GICA', 'organo': TRANSVERSAL},
    '6': {'tema': 'Modernización de TI', 'organo': 'SG'},
    '7': {'tema': 'Certificación ambiental', 'organo': 'SENACE'},
    '8': {'tema': 'Optimización presupuestal y financiera', 'organo': 'OGPP'},
    '9': {'tema': 'Economías ilegales', 'organo': TRANSVERSAL},
    '10': {'tema': 'Fiscalización ambiental', 'organo': 'OEFA'},
    '11 OII': {'tema': 'Integridad pública', 'organo': 'SG'},
    '11 RRHH': {'tema': 'Gestión de recursos humanos', 'organo': 'SG'},
}

NOISE = re.compile(r'^(?:[ÁA]RBOL DE PROBLEMA|MESA DE TRABAJO|Nota:)', re.IGNORECASE)
NOTE_BOX = re.compile(r'^(\(?\*+\)?)\s*(.+)$', re.S)

CODE_PATTERNS = [
    # (regex, lado) — el grupo 'n' es el código numérico jerárquico
    (re.compile(r'^PROBLEMA CENTRAL\s*\|?\s*', re.I), 'central'),
    (re.compile(r'^EFECTO (?:FINAL|INDIRECTO|DIRECTO)\s*(?P<n>\d+(?:\.\d+)*)?\s*\|?\s*', re.I), 'effect'),
    (re.compile(r'^CAUSA (?:DIRECTA|INDIRECTA)\s*(?P<n>\d+(?:\.\d+)*)?\s*\|?\s*', re.I), 'cause'),
    (re.compile(r'^(?:EFF|EF|EI|ED)\s*-?\s*(?P<n>\d+(?:\.\d+)*)\.?\s*\|?\s*'), 'effect'),
    (re.compile(r'^(?:CD|CI|C)\s*-?\s*(?P<n>\d+(?:\.\d+)*)\.?\s*\|?\s*'), 'cause'),
    (re.compile(r'^(?P<n>\d+(?:\.\d+)*)\.?\s+(?=[A-ZÁÉÍÓÚÑ«“"])'), None),
]


@dataclass
class RawNode:
    key: str
    text: str
    side: str                      # 'central' | 'cause' | 'effect'
    x: float = 0.0
    y: float = 0.0
    w: float = 0.0
    h: float = 0.0
    code: str | None = None        # código jerárquico numérico ("2.3")
    code_label: str | None = None  # prefijo original ("CI 2.3")
    style: tuple = ()
    source: str = ''
    notes: list[str] = field(default_factory=list)
    parents: list[str] = field(default_factory=list)
    method: str = ''
    shape_id: str | None = None

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def cx(self) -> float:
        return self.x + self.w / 2


# ------------------------------------------------------------------ utilidades de texto

def clean_text(t: str) -> str:
    t = t.replace(' ', ' ').replace('​', '')
    t = re.sub(r'[ \t]+', ' ', t)
    t = re.sub(r'\s*\n\s*', '\n', t).strip()
    return t


def split_code(text: str) -> tuple[str | None, str | None, str | None, str]:
    """Devuelve (código numérico, prefijo original, lado sugerido, texto sin prefijo)."""
    one_line = re.sub(r'\s*\n\s*', ' | ', text, count=1) if '\n' in text else text
    for rx, side in CODE_PATTERNS:
        m = rx.match(one_line)
        if m:
            body = one_line[m.end():].strip()
            body = re.sub(r'\s*\n\s*', ' ', body.replace(' | ', ' ')).strip()
            n = m.groupdict().get('n')
            return n, m.group(0).strip(' |'), side, body
    return None, None, None, re.sub(r'\s*\n\s*', ' ', text).strip()


def mesa_key(path: Path) -> str:
    name = unicodedata.normalize('NFC', path.stem)
    m = re.search(r'Mesa (\d+)', name)
    num = m.group(1)
    if num == '11':
        return '11 OII' if 'OII' in name else '11 RRHH'
    return num


def slug(key: str) -> str:
    return 'm' + key.lower().replace(' ', '')


# ------------------------------------------------------------------ geometría

def rect_dist(p: tuple[float, float], s) -> float:
    dx = max(s.x - p[0], 0, p[0] - (s.x + s.w))
    dy = max(s.y - p[1], 0, p[1] - (s.y + s.h))
    return (dx * dx + dy * dy) ** 0.5


def point_seg_dist(p, a, b) -> float:
    (px, py), (ax, ay), (bx, by) = p, a, b
    vx, vy = bx - ax, by - ay
    L = vx * vx + vy * vy
    t = 0.0 if L == 0 else max(0.0, min(1.0, ((px - ax) * vx + (py - ay) * vy) / L))
    cx, cy = ax + t * vx, ay + t * vy
    return ((px - cx) ** 2 + (py - cy) ** 2) ** 0.5


def on_path(p, path) -> bool:
    return any(point_seg_dist(p, path[i], path[i + 1]) <= TOUCH for i in range(len(path) - 1))


class DSU:
    def __init__(self, n: int):
        self.p = list(range(n))

    def find(self, i: int) -> int:
        while self.p[i] != i:
            self.p[i] = self.p[self.p[i]]
            i = self.p[i]
        return i

    def union(self, a: int, b: int) -> None:
        self.p[self.find(a)] = self.find(b)


# ------------------------------------------------------------------ cajas → nodos

def nodes_from_shapes(shapes: list[Shape], source_file: str) -> tuple[list[RawNode], list[Shape], dict[str, str]]:
    notes: dict[str, str] = {}
    boxes: list[Shape] = []
    for s in shapes:
        if s.connector or not s.text.strip() or s.w <= 0 or s.h <= 0:
            continue
        t = clean_text(s.text)
        if NOISE.match(t):
            continue
        nm = NOTE_BOX.match(t)
        if nm and (t.startswith('(*') or t.startswith('*')):
            notes[nm.group(1).strip('()')] = nm.group(2).strip()
            continue
        # duplicados (misma caja dibujada dos veces)
        if any(b.text == s.text and abs(b.x - s.x) < CM and abs(b.y - s.y) < CM for b in boxes):
            continue
        boxes.append(s)
    lines = [s for s in shapes if s.connector]

    raw: list[RawNode] = []
    for i, b in enumerate(boxes):
        # una caja con lista numerada automática contiene varias causas/efectos hermanos
        items = b.extra.get('items') or [b.text]
        for k, item in enumerate(items):
            code, code_label, side_hint, body = split_code(clean_text(item))
            if len(items) > 1:
                code, code_label = None, f'ítem {k + 1}'
            raw.append(RawNode(key=f'b{i}' + (f'_{k}' if len(items) > 1 else ''), text=body, side=side_hint or '',
                               x=b.x, y=b.y + k, w=b.w, h=b.h, code=code, code_label=code_label,
                               style=(b.fill.upper(), b.bold), source=f'{source_file} · {b.source}',
                               notes=list(b.footnote_marks), shape_id=b.extra.get('id') if k == 0 else None))
    return raw, lines, notes


def find_central(nodes: list[RawNode]) -> RawNode:
    tagged = [n for n in nodes if n.side == 'central']
    if tagged:
        return tagged[0]
    maxw = max(n.w for n in nodes)
    cands = [n for n in nodes if n.w >= 0.75 * maxw]

    def score(c: RawNode) -> tuple[int, float]:
        above = sum(o.cy < c.y for o in nodes)
        below = sum(o.cy > c.y + c.h for o in nodes)
        return min(above, below), c.w
    return max(cands, key=score)


def assign_sides(nodes: list[RawNode], central: RawNode) -> None:
    for n in nodes:
        if n is central:
            n.side = 'central'
        elif n.cy < central.cy:
            n.side = 'effect'
        else:
            n.side = 'cause'


def closeness(n: RawNode, side: str) -> float:
    """Distancia al problema central (menor = más cercano) medida por el borde próximo."""
    return n.y if side == 'cause' else -(n.y + n.h)


def link_by_codes(nodes: list[RawNode]) -> None:
    by_code: dict[tuple[str, str], RawNode] = {}
    repeated: set[tuple[str, str]] = set()
    for n in nodes:
        if n.code:
            k = (n.side, n.code)
            if k in by_code:
                repeated.add(k)
            by_code.setdefault(k, n)
    for k in repeated:  # numeraciones locales que se repiten (p. ej. "1." en cada columna) no son jerárquicas
        del by_code[k]
    for n in nodes:
        if n.code and '.' in n.code:
            parent = by_code.get((n.side, n.code.rsplit('.', 1)[0]))
            if parent is not None and parent is not n:
                n.parents = [parent.key]
                n.method = 'código'


def link_by_connectors(nodes: list[RawNode], lines: list[Shape], central: RawNode, shape_ids: dict) -> None:
    if not lines:
        return
    paths = [ln.path for ln in lines]
    dsu = DSU(len(lines))
    for i, ln in enumerate(lines):
        for p in ln.endpoints:
            for j in range(len(lines)):
                if i != j and on_path(p, paths[j]):
                    dsu.union(i, j)
    touch: dict[int, set[str]] = {}
    for i, ln in enumerate(lines):
        comp = dsu.find(i)
        for p in ln.endpoints:
            for n in nodes:
                if rect_dist(p, n) <= TOUCH:
                    touch.setdefault(comp, set()).add(n.key)
        for ref in (ln.extra.get('st'), ln.extra.get('en')):
            if ref and ref in shape_ids:
                touch.setdefault(comp, set()).add(shape_ids[ref])
    by_key = {n.key: n for n in nodes}
    for members in touch.values():
        group = [by_key[k] for k in members]
        for side in ('cause', 'effect'):
            side_nodes = [n for n in group if n.side == side]
            if not side_nodes:
                continue
            pool = side_nodes + ([central] if central in group else [])
            best = min(closeness(n, side) for n in pool)
            parents = [n for n in pool if closeness(n, side) <= best + ROW_TOL]
            for child in pool:
                if child in parents or child.method == 'código':
                    continue
                for p in parents:
                    if p.key not in child.parents:
                        child.parents.append(p.key)
                child.method = 'conector'


def link_by_geometry(nodes: list[RawNode], central: RawNode) -> None:
    for n in nodes:
        if n.side == 'central' or n.parents:
            continue
        cands = []
        for c in nodes:
            if c is n or (c.side != n.side and c is not central):
                continue
            nearer = (c.y + c.h <= n.y + ROW_TOL) if n.side == 'cause' else (c.y >= n.y + n.h - ROW_TOL)
            overlap = min(c.x + c.w, n.x + n.w) - max(c.x, n.x)
            if nearer and overlap > 0 and c.style != n.style:
                gap = (n.y - (c.y + c.h)) if n.side == 'cause' else (c.y - (n.y + n.h))
                cands.append((gap, -overlap, c))
        parent = min(cands, key=lambda t: (t[0], t[1]))[2] if cands else central
        n.parents = [parent.key]
        n.method = 'geometría'


def prune_and_level(nodes: list[RawNode], central: RawNode) -> dict[str, int]:
    """Niveles por camino más largo; elimina atajos y aristas que apuntan alejándose del centro."""
    by_key = {n.key: n for n in nodes}
    level: dict[str, int] = {central.key: 0}
    for n in nodes:  # descartar padres del semiplano opuesto
        n.parents = [p for p in n.parents if by_key[p].side in (n.side, 'central') and p != n.key]
        if n.side != 'central' and not n.parents:
            n.parents = [central.key]
            n.method = n.method or 'por defecto'

    def depth(k: str, stack: frozenset) -> int:
        if k in level:
            return abs(level[k])
        node = by_key[k]
        valid = [p for p in node.parents if p not in stack]
        if not valid:
            valid = [central.key]
        d = 1 + max(depth(p, stack | {k}) for p in valid)
        node.parents = valid
        level[k] = d if node.side == 'cause' else -d
        return d

    for n in nodes:
        depth(n.key, frozenset())
    for n in nodes:  # conservar solo padres del nivel inmediato (quita atajos)
        if n.side == 'central':
            continue
        target = abs(level[n.key]) - 1
        n.parents = [p for p in n.parents if abs(level[p]) == target] or n.parents
    return level


# ------------------------------------------------------------------ Mesa 10 (tabla)

def nodes_from_table(path: Path, source_file: str) -> list[RawNode]:
    d = docx.Document(str(path))
    tbl = d.tables[0]
    nodes: list[RawNode] = []
    rows: dict[str, list[RawNode]] = {}
    for r_i, tr in enumerate(tbl._tbl.tr_lst):
        col = 0
        cells = []
        for tc in tr.tc_lst:
            span = tc.grid_span
            cells.append((col, span, tc))
            col += span
        header = ''.join(t.text or '' for t in cells[0][2].iter(qn('w:t'))).strip().upper()
        if not header:
            continue
        for c0, span, tc in cells[1:]:
            paras = [paragraph_text(p).strip() for p in tc.iter(qn('w:p'))]
            paras = [p for p in paras if p]
            entries = paras if header.startswith('SUBCAUSA') else [' '.join(paras)]
            for k, txt in enumerate(entries):
                code, code_label, side, body = split_code(txt)
                side = {'EFECTOS FINALES': 'effect', 'EFECTOS DIRECTOS': 'effect',
                        'PROBLEMA CENTRAL': 'central'}.get(header, 'cause')
                if header.startswith('SUBCAUSA') and code is None and nodes:
                    # primera subcausa sin numeración explícita: hereda la columna
                    col_parent = next(n for n in rows['CAUSAS DIRECTAS'] if n.x <= c0 < n.x + n.w)
                    code = f'{col_parent.code}.1'
                    code_label = code
                n = RawNode(key=f't{r_i}_{c0}_{k}', text=body, side=side, x=c0, y=r_i * 10 + k, w=span, h=1,
                            code=code, code_label=code_label, source=f'{source_file} · tabla, fila «{header.title()}»')
                rows.setdefault(header, []).append(n)
                nodes.append(n)
    central = rows['PROBLEMA CENTRAL'][0]
    for n in rows['EFECTOS DIRECTOS'] + rows['CAUSAS DIRECTAS']:
        n.parents, n.method = [central.key], 'tabla'
    for f in rows['EFECTOS FINALES']:  # efectos finales: hijos de los efectos directos con solape de columnas
        f.parents = [e.key for e in rows['EFECTOS DIRECTOS'] if min(e.x + e.w, f.x + f.w) - max(e.x, f.x) > 0]
        f.method = 'tabla'
    link_by_codes(nodes)
    return nodes


def table_footnotes(path: Path) -> dict[str, str]:
    d = docx.Document(str(path))
    out = {}
    for p in d.paragraphs:
        m = re.match(r'^\((\d+)\)\s*(.+)', p.text.strip())
        if m:
            out[m.group(1)] = m.group(2)
    return out


def body_footnotes(path: Path) -> dict[str, str]:
    """Notas al pie numeradas en el cuerpo del documento (listas automáticas o '1. texto')."""
    d = docx.Document(str(path))
    out: dict[str, str] = {}
    i = 0
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        is_list = p._p.pPr is not None and p._p.pPr.numPr is not None
        m = re.match(r'^\(?(\d+)[.)]\s+(.+)', t)
        if m:
            out[m.group(1)] = m.group(2)
        elif is_list:
            i += 1
            out[str(i)] = t
    return out


# ------------------------------------------------------------------ construcción del árbol

def build_tree(key: str, src: Path, raw: list[RawNode], central: RawNode, level: dict[str, int],
               notes: dict[str, str], footnotes: dict[str, str]) -> dict:
    meta = MESAS[key]
    prefix = slug(key)
    by_key = {n.key: n for n in raw}
    children: dict[str, list[RawNode]] = {}
    for n in raw:
        for p in n.parents:
            children.setdefault(p, []).append(n)

    def order(ns: list[RawNode]) -> list[RawNode]:
        return sorted(ns, key=lambda n: (round(n.x / (2 * CM)) if n.w > 50 else n.x, n.y))

    ids: dict[str, str] = {central.key: f'{prefix}-pc'}
    queue = [central]
    while queue:
        cur = queue.pop(0)
        for side, letter in (('cause', 'c'), ('effect', 'e')):
            kids = order([k for k in children.get(cur.key, []) if k.side == side and k.key not in ids])
            for i, k in enumerate(kids, start=1):
                base = ids[cur.key]
                ids[k.key] = f'{prefix}-{letter}{i}' if cur is central else f'{base}-{i}'
                queue.append(k)

    descendants: dict[str, int] = {}

    def count(k: str, seen: set) -> set:
        for c in children.get(k, []):
            if c.key not in seen:
                seen.add(c.key)
                count(c.key, seen)
        return seen
    for n in raw:
        descendants[ids[n.key]] = len(count(n.key, set()))

    out_nodes = []
    for n in sorted(raw, key=lambda n: (level[n.key], ids[n.key])):
        lv = level[n.key]
        ntype = ('problema_central' if lv == 0 else 'causa_directa' if lv == 1 else 'causa_indirecta' if lv > 1
                 else 'efecto_directo' if lv == -1 else 'efecto_indirecto')
        text = n.text
        node_notes = []
        for mark in n.notes:
            if mark in footnotes:
                node_notes.append(f'Nota {mark}: {footnotes[mark]}')
        m_paren = re.findall(r'\((\d)\)', text)
        for mark in m_paren:
            if mark in footnotes:
                node_notes.append(f'Nota ({mark}): {footnotes[mark]}')
        for mark, note in sorted(notes.items(), key=lambda kv: -len(kv[0])):
            if re.search(r'(?<!\*)' + re.escape(mark) + r'(?!\*)', text.replace('(', '').replace(')', '')):
                node_notes.append(f'({mark}) {note}')
        fuente = n.source + (f' · {n.code_label.rstrip(".")}' if n.code_label else '')
        node = {
            'id': ids[n.key],
            'type': ntype,
            'level': lv,
            'label': make_label(text),
            'description': text,
            'organo': classify_organo(text, meta['organo']),
            'ejeTematico': classify_eje(text, n.side, meta['tema']),
            'severidad': 'moderada',
            'severidad_manual': False,
            'fuente': fuente,
        }
        ev = extract_evidencia(text, node_notes)
        if ev:
            node['evidencia'] = ev
        node['_method'] = n.method
        out_nodes.append(node)
    assign_severity(out_nodes, descendants)

    edges = []
    for n in raw:
        for p in n.parents:
            if n.side == 'cause':
                s, t = ids[n.key], ids[p]
            else:
                s, t = ids[p], ids[n.key]
            edges.append({'id': f'{s}__{t}', 'source': s, 'target': t})
    edges.sort(key=lambda e: e['id'])

    counts = {t: sum(1 for n in out_nodes if n['type'] == t) for t in
              ('causa_directa', 'causa_indirecta', 'efecto_directo', 'efecto_indirecto')}
    titulo_mesa = f'Mesa {key.split()[0]}' + (f' · {key.split()[1]}' if ' ' in key else '')
    return {
        'id': prefix,
        'titulo': f'{titulo_mesa} — {meta["tema"]}',
        'descripcionGeneral': (
            f'Árbol de problemas de la {titulo_mesa} ({meta["tema"]}). Problema central: '
            f'«{central.text.rstrip(".")}». Identifica {counts["causa_directa"]} causas directas, '
            f'{counts["causa_indirecta"]} causas indirectas, {counts["efecto_directo"]} efectos directos y '
            f'{counts["efecto_indirecto"]} efectos indirectos.'),
        'archivoOrigen': unicodedata.normalize('NFC', src.name),
        'problemaCentralId': ids[central.key],
        'nodes': out_nodes,
        'edges': edges,
    }


def process(src: Path) -> dict:
    key = mesa_key(src)
    fname = unicodedata.normalize('NFC', src.name)
    notes: dict[str, str] = {}
    footnotes: dict[str, str] = {}
    if src.suffix.lower() == '.pptx':
        slides = pptx_slides(src)
        num, shapes = max(slides, key=lambda s: sum(1 for x in s[1] if x.text.strip()))
        raw, lines, notes = nodes_from_shapes(shapes, fname)
        shape_ids = {n.shape_id: n.key for n in raw if n.shape_id}
    else:
        norm = NORMALIZED / (unicodedata.normalize('NFC', src.stem) + '.docx')
        path = norm if norm.exists() else src
        if path is src and src.suffix.lower() == '.doc':
            raise SystemExit(f'Falta la versión normalizada de {fname}: ejecute scripts/normalize_word.py')
        shapes = docx_shapes(path)
        footnotes = body_footnotes(path)
        if not [s for s in shapes if s.text.strip()]:
            raw = nodes_from_table(path, fname)
            footnotes = table_footnotes(path) or footnotes
            central = next(n for n in raw if n.side == 'central')
            level = prune_and_level(raw, central)
            return build_tree(key, src, raw, central, level, notes, footnotes)
        raw, lines, notes = nodes_from_shapes(shapes, fname)
        shape_ids = {}
    central = find_central(raw)
    assign_sides(raw, central)
    link_by_codes(raw)
    link_by_connectors(raw, lines, central, shape_ids)
    link_by_geometry(raw, central)
    level = prune_and_level(raw, central)
    return build_tree(key, src, raw, central, level, notes, footnotes)


def mesa_sort_key(p: Path) -> tuple:
    k = mesa_key(p)
    return int(k.split()[0]), k


def report(trees: list[dict]) -> str:
    lines = []
    for t in trees:
        lines.append(f'\n## {t["titulo"]}  ({t["archivoOrigen"]})\n')
        nodes = {n['id']: n for n in t['nodes']}
        kids: dict[str, list[str]] = {}
        for e in t['edges']:
            parent, child = (e['target'], e['source']) if nodes[e['source']]['level'] > 0 else (e['source'], e['target'])
            kids.setdefault(parent, []).append(child)

        def walk(nid: str, depth: int, side: int):
            n = nodes[nid]
            lines.append(f'{"  " * depth}- [{n["id"]}] L{n["level"]} {n["severidad"][:4]} {n["organo"]} '
                         f'({n["_method"]}) {n["label"]}')
            for c in sorted(kids.get(nid, [])):
                if (nodes[c]['level'] > 0) == (side > 0):
                    walk(c, depth + 1, side)
        lines.append('EFECTOS')
        walk(t['problemaCentralId'], 0, -1)
        lines.append('CAUSAS')
        walk(t['problemaCentralId'], 0, 1)
    return '\n'.join(lines)


def preserve_manual_severity(trees: list[dict], previous_path: Path) -> list[str]:
    """Respeta las severidades corregidas a mano (`severidad_manual: true`) del JSON anterior.

    El nodo se empareja por su texto íntegro dentro del mismo árbol (estable aunque cambie la
    numeración de IDs) y, en su defecto, por ID con el mismo texto. Devuelve avisos de las
    correcciones manuales que ya no encuentran su nodo.
    """
    if not previous_path.exists():
        return []
    prev = json.loads(previous_path.read_text(encoding='utf-8'))
    warnings = []
    current = {t['id']: t for t in trees}
    for pt in prev.get('arboles', []):
        tree = current.get(pt['id'])
        for pn in pt.get('nodes', []):
            if not pn.get('severidad_manual'):
                continue
            match = None
            if tree:
                match = next((n for n in tree['nodes'] if n['description'] == pn['description']), None)
            if match is None:
                warnings.append(f'{pt["id"]}/{pn["id"]}: severidad manual sin nodo equivalente '
                                f'(«{pn["label"]}»); revise la corrección')
                continue
            match['severidad'] = pn['severidad']
            match['severidad_manual'] = True
    return warnings


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--report', type=Path)
    args = ap.parse_args()
    files = sorted((p for p in RAW.iterdir() if p.suffix.lower() in ('.doc', '.docx', '.pptx')), key=mesa_sort_key)
    trees = [process(f) for f in files]
    if args.report:
        args.report.write_text(report(trees), encoding='utf-8')
    for t in trees:
        for n in t['nodes']:
            n.pop('_method', None)
    for w in preserve_manual_severity(trees, OUT):
        print('AVISO:', w)
    manual = sum(n['severidad_manual'] for t in trees for n in t['nodes'])
    if manual:
        print(f'{manual} severidad(es) manual(es) preservada(s)')
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({'arboles': trees}, ensure_ascii=False, indent=2), encoding='utf-8')
    organos = {sigla: {'nombre': nombre, 'dependencia': dep} for sigla, (nombre, dep) in ORGANOS.items()}
    OUT_ORGANOS.write_text(json.dumps(organos, ensure_ascii=False, indent=2), encoding='utf-8')
    total = sum(len(t['nodes']) for t in trees)
    print(f'{len(trees)} árboles, {total} nodos → {OUT.relative_to(ROOT)}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
